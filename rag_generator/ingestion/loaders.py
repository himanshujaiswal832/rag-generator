"""
Universal Document Loaders for RAG Generator
Supports PDF, DOCX, TXT, Markdown, CSV, TSV, JSON, HTML, and Raw Text
"""

import os
import csv
import json
import hashlib
from typing import List, Dict, Any, Union, BinaryIO
from io import BytesIO, StringIO
from bs4 import BeautifulSoup
import pypdf
import docx

from rag_generator.core.models import Document


class DocumentLoader:
    """Universal Loader that automatically detects file format and parses content."""

    SUPPORTED_EXTENSIONS = {
        ".pdf": "pdf",
        ".docx": "docx",
        ".doc": "docx",
        ".txt": "text",
        ".md": "markdown",
        ".markdown": "markdown",
        ".csv": "csv",
        ".tsv": "tsv",
        ".json": "json",
        ".jsonl": "jsonl",
        ".html": "html",
        ".htm": "html",
    }

    @staticmethod
    def compute_sha256(content: Union[str, bytes]) -> str:
        """Compute SHA256 checksum of document content."""
        if isinstance(content, str):
            content = content.encode("utf-8", errors="ignore")
        return hashlib.sha256(content).hexdigest()

    @classmethod
    def load_from_file(cls, file_path: str) -> Document:
        """Load a single document from a filesystem path."""
        if not os.path.exists(file_path):
            raise FileNotFoundError(f"File not found: {file_path}")

        filename = os.path.basename(file_path)
        _, ext = os.path.splitext(filename.lower())
        file_type = cls.SUPPORTED_EXTENSIONS.get(ext, "text")

        with open(file_path, "rb") as f:
            file_bytes = f.read()

        return cls.load_from_bytes(file_bytes, filename=filename, file_type=file_type)

    @classmethod
    def load_from_bytes(
        cls,
        file_bytes: bytes,
        filename: str,
        file_type: str = "text",
    ) -> Document:
        """Parse bytes into a structured Document object."""
        _, ext = os.path.splitext(filename.lower())
        if ext in cls.SUPPORTED_EXTENSIONS:
            file_type = cls.SUPPORTED_EXTENSIONS[ext]

        sha256 = cls.compute_sha256(file_bytes)
        file_size = len(file_bytes)

        if file_type == "pdf":
            content, pages = cls._parse_pdf(file_bytes)
        elif file_type == "docx":
            content, pages = cls._parse_docx(file_bytes)
        elif file_type == "csv":
            content, pages = cls._parse_csv(file_bytes, delimiter=",")
        elif file_type == "tsv":
            content, pages = cls._parse_csv(file_bytes, delimiter="\t")
        elif file_type in ("json", "jsonl"):
            content, pages = cls._parse_json(file_bytes, is_jsonl=(file_type == "jsonl"))
        elif file_type == "html":
            content, pages = cls._parse_html(file_bytes)
        else:
            # Default plain text / markdown
            content, pages = cls._parse_text(file_bytes)

        metadata = {
            "file_size": file_size,
            "sha256": sha256,
            "page_count": len(pages) if pages else 1,
            "original_filename": filename,
        }

        return Document(
            filename=filename,
            file_type=file_type,
            content=content,
            metadata=metadata,
            pages=pages,
        )

    @classmethod
    def load_from_text(cls, text: str, title: str = "Pasted_Text.txt") -> Document:
        """Create a Document object directly from raw text content."""
        sha256 = cls.compute_sha256(text)
        return Document(
            filename=title,
            file_type="text",
            content=text,
            metadata={
                "file_size": len(text.encode("utf-8")),
                "sha256": sha256,
                "page_count": 1,
            },
            pages=[{"page_number": 1, "text": text}],
        )

    @staticmethod
    def _parse_pdf(file_bytes: bytes):
        """Extract text and page-by-page mapping from PDF."""
        stream = BytesIO(file_bytes)
        reader = pypdf.PdfReader(stream)
        pages_data = []
        full_text_parts = []

        for idx, page in enumerate(reader.pages):
            page_text = page.extract_text() or ""
            page_num = idx + 1
            if page_text.strip():
                pages_data.append({"page_number": page_num, "text": page_text})
                full_text_parts.append(f"--- Page {page_num} ---\n{page_text}")

        full_content = "\n\n".join(full_text_parts)
        if not full_content.strip():
            full_content = "[PDF with no extractable text or scanned images]"
            pages_data = [{"page_number": 1, "text": full_content}]

        return full_content, pages_data

    @staticmethod
    def _parse_docx(file_bytes: bytes):
        """Extract text and sections from Word docx."""
        stream = BytesIO(file_bytes)
        doc = docx.Document(stream)
        paragraphs = []
        for p in doc.paragraphs:
            if p.text.strip():
                paragraphs.append(p.text.strip())

        # Also extract text from tables
        for table in doc.tables:
            table_rows = []
            for row in table.rows:
                cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                if cells:
                    table_rows.append(" | ".join(cells))
            if table_rows:
                paragraphs.append("\n".join(table_rows))

        full_content = "\n\n".join(paragraphs)
        pages_data = [{"page_number": 1, "text": full_content}]
        return full_content, pages_data

    @staticmethod
    def _parse_csv(file_bytes: bytes, delimiter: str = ","):
        """Extract structured representation from CSV/TSV."""
        text = file_bytes.decode("utf-8", errors="replace")
        reader = csv.reader(StringIO(text), delimiter=delimiter)
        rows = list(reader)
        if not rows:
            return "", []

        headers = rows[0]
        data_rows = rows[1:]

        formatted_lines = [f"Table Headers: {', '.join(headers)}"]
        for idx, row in enumerate(data_rows, 1):
            row_items = []
            for h, val in zip(headers, row):
                row_items.append(f"{h}: {val}")
            # If row has more columns than headers
            if len(row) > len(headers):
                for val in row[len(headers):]:
                    row_items.append(f"Extra: {val}")
            formatted_lines.append(f"Row {idx} -> " + ", ".join(row_items))

        full_content = "\n".join(formatted_lines)
        pages_data = [{"page_number": 1, "text": full_content}]
        return full_content, pages_data

    @staticmethod
    def _parse_json(file_bytes: bytes, is_jsonl: bool = False):
        """Extract readable text from JSON or JSONL."""
        text = file_bytes.decode("utf-8", errors="replace")
        if is_jsonl:
            lines = [l.strip() for l in text.splitlines() if l.strip()]
            records = []
            for l in lines:
                try:
                    records.append(json.loads(l))
                except Exception:
                    records.append({"raw_line": l})
            formatted = json.dumps(records, indent=2)
        else:
            try:
                parsed = json.loads(text)
                formatted = json.dumps(parsed, indent=2)
            except Exception:
                formatted = text

        pages_data = [{"page_number": 1, "text": formatted}]
        return formatted, pages_data

    @staticmethod
    def _parse_html(file_bytes: bytes):
        """Extract clean text hierarchy from HTML."""
        html_text = file_bytes.decode("utf-8", errors="replace")
        soup = BeautifulSoup(html_text, "html.parser")

        # Strip scripts, styles, noscript, svg
        for tag in soup(["script", "style", "noscript", "svg", "header", "footer", "nav"]):
            tag.decompose()

        title = soup.title.string.strip() if soup.title and soup.title.string else "Document"
        text = soup.get_text(separator="\n", strip=True)
        content = f"Title: {title}\n\n{text}"
        pages_data = [{"page_number": 1, "text": content}]
        return content, pages_data

    @staticmethod
    def _parse_text(file_bytes: bytes):
        """Extract plain text/markdown."""
        text = file_bytes.decode("utf-8", errors="replace")
        pages_data = [{"page_number": 1, "text": text}]
        return text, pages_data
